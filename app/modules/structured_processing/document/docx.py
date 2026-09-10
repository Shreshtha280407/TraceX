"""DOCX text extraction via `python-docx`.

Extracts body paragraphs and text-bearing table cells into a single,
deterministic text stream (paragraphs first, then tables, both in their
own document order). This does not preserve true interleaved visual
order for a table embedded mid-document — see `docs/qa/known-limitations.md`
— but it is fully deterministic: the same DOCX always yields the same
stream and the same span offsets, which is what downstream provenance
needs. `python-docx` never executes macros or follows external links; it
only reads the document XML.
"""

from __future__ import annotations

import io

from docx import Document

from app.modules.structured_processing.document.text_extractors import enforce_text_limits
from app.modules.structured_processing.errors import ErrorCode, ProcessingError
from app.modules.structured_processing.limits import validate_zip_container
from app.modules.structured_processing.models import TextSegment


def extract_docx(data: bytes) -> list[TextSegment]:
    """Extract a DOCX into a single `TextSegment` (`page=None`).

    Raises `invalid_docx` for a corrupt file or one that fails the ZIP
    container safety check, and `input_limit_exceeded` if the combined
    text is too large.
    """
    validate_zip_container(data, error_code=ErrorCode.INVALID_DOCX)

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises varied errors on bad input
        raise ProcessingError(ErrorCode.INVALID_DOCX, "DOCX could not be parsed") from exc

    pieces: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            pieces.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    pieces.append(text)

    combined = "\n".join(pieces)
    segments = [TextSegment(page=None, text=combined)]
    enforce_text_limits(segments)
    return segments
