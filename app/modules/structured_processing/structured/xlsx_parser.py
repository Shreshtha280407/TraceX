"""Safe XLSX parsing via `openpyxl`.

Opened with `data_only=False` (the default) so a formula cell's `.value`
is the formula *string itself* (e.g. `"=SUM(A1:A2)"`) — openpyxl never
evaluates formulas either way, but `data_only=False` also avoids silently
trusting a workbook's last-cached computed value, which could be stale or
have been hand-edited independently of the formula. Either way, a formula
cell's content is treated as plain source text, never as an instruction.
"""

from __future__ import annotations

import io
from collections.abc import Callable

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.contracts.common import SourceLocator
from app.modules.structured_processing.errors import ErrorCode, ProcessingError
from app.modules.structured_processing.limits import (
    MAX_CELL_TEXT_LENGTH,
    MAX_COLUMNS,
    MAX_ROWS,
    validate_zip_container,
)
from app.modules.structured_processing.models import RawRecord


def parse_xlsx(data: bytes, *, sheet_name: str | None = None) -> list[RawRecord]:
    """Parse one XLSX sheet into `RawRecord`s (defaults to the first sheet).

    Raises `invalid_xlsx` for a corrupt file, a failed ZIP container
    safety check, a missing header row, or an unknown `sheet_name`; raises
    `input_limit_exceeded` if rows, columns, or a cell's text exceed the
    documented limits.
    """
    validate_zip_container(data, error_code=ErrorCode.INVALID_XLSX)

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (InvalidFileException, KeyError, OSError) as exc:
        raise ProcessingError(ErrorCode.INVALID_XLSX, "XLSX could not be parsed") from exc

    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    except KeyError as exc:
        raise ProcessingError(ErrorCode.INVALID_XLSX, f"sheet '{sheet_name}' not found") from exc
    except IndexError as exc:
        raise ProcessingError(ErrorCode.INVALID_XLSX, "workbook has no sheets") from exc

    rows_iter = sheet.iter_rows(values_only=False)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ProcessingError(ErrorCode.INVALID_XLSX, "sheet has no header row") from exc

    if len(header_row) > MAX_COLUMNS:
        raise ProcessingError(
            ErrorCode.INPUT_LIMIT_EXCEEDED, f"sheet exceeds the {MAX_COLUMNS}-column limit"
        )

    column_index: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = str(cell.value).strip() if cell.value is not None else ""
        if name and name not in column_index:
            column_index[name] = idx

    records: list[RawRecord] = []
    for row_number, row in enumerate(rows_iter, start=2):  # row 1 is the header
        data_row_index = row_number - 2
        if data_row_index >= MAX_ROWS:
            raise ProcessingError(
                ErrorCode.INPUT_LIMIT_EXCEEDED, f"sheet exceeds the {MAX_ROWS}-row limit"
            )
        if len(row) > MAX_COLUMNS:
            raise ProcessingError(
                ErrorCode.INPUT_LIMIT_EXCEEDED,
                f"sheet row {row_number} exceeds the {MAX_COLUMNS}-column limit",
            )
        if all(cell.value is None for cell in row):
            continue  # fully empty row: nothing to observe

        values: dict[str, str] = {}
        for name, idx in column_index.items():
            if idx >= len(row):
                continue
            raw_value = row[idx].value
            if raw_value is None:
                continue
            text_value = str(raw_value)
            if len(text_value) > MAX_CELL_TEXT_LENGTH:
                raise ProcessingError(
                    ErrorCode.INPUT_LIMIT_EXCEEDED,
                    f"sheet row {row_number} has a cell exceeding "
                    f"{MAX_CELL_TEXT_LENGTH} characters",
                )
            values[name] = text_value

        records.append(
            RawRecord(
                index=data_row_index,
                values=values,
                locator_for=_make_locator(sheet.title, row_number, column_index),
            )
        )

    workbook.close()
    return records


def _make_locator(
    sheet_title: str, row_number: int, column_index: dict[str, int]
) -> Callable[[str | None], SourceLocator]:
    def locator_for(field_name: str | None) -> SourceLocator:
        if field_name is None:
            return SourceLocator(sheet=sheet_title, row=row_number)
        idx = column_index.get(field_name)
        if idx is None:
            return SourceLocator(sheet=sheet_title, row=row_number)
        return SourceLocator(sheet=sheet_title, row=row_number, column=idx + 1)

    return locator_for
