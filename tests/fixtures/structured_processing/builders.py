"""Synthetic, non-sensitive fixture builders for structured-processing tests.

No `reportlab` (not an allowed dependency for this phase): PDFs are built
by hand-assembling a minimal, valid single-font PDF byte stream directly,
which is enough to exercise real `pypdf` extraction without adding a new
dependency. DOCX/XLSX fixtures use `python-docx`/`openpyxl` directly, the
same libraries the app code parses them with.

`build_scanned_pdf_page` (Phase 3 -- Jasraj) hand-assembles a PDF page
whose only content is a raster image XObject (real, Pillow-rendered text,
no embedded text layer at all) -- for real Tesseract OCR fallback tests.
Uses the same "no new PDF-generation dependency" hand-assembly technique
`build_minimal_pdf` already established, extended to embed an uncompressed
`/DeviceRGB` image stream (a real, valid PDF image XObject; no filter is a
legal, if uncompressed, way to store raw samples).
"""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
)


def _load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in _FONT_CANDIDATES:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def build_scanned_pdf_page(
    lines: list[str], *, width: int = 900, height: int | None = None
) -> bytes:
    """A single-page PDF whose only content is a rendered image of `lines` -- no text layer.

    Each line is rendered at a fixed vertical offset via Pillow (the same
    "render text with whatever system font is available" technique
    `tests/unit/media_processing/test_tesseract_ocr.py` already uses for
    Gaurav's OCR tests), so `pypdf.extract_text()` on the resulting PDF
    always returns an empty string while real Tesseract OCR against a
    rendered page genuinely recovers the text -- a real scanned-document
    fixture, not a simulated one.
    """
    line_height = 60
    resolved_height = height or (line_height * len(lines) + 40)
    image = Image.new("RGB", (width, resolved_height), color=(255, 255, 255))
    draw = ImageDraw.Draw(image)
    font = _load_font(32)
    for index, line in enumerate(lines):
        draw.text((20, 20 + index * line_height), line, fill=(0, 0, 0), font=font)

    raw_rgb = image.tobytes()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            f"<< /Type /Page /Parent 2 0 R /Resources << /XObject << /Im0 5 0 R >> >> "
            f"/MediaBox [0 0 {image.width} {image.height}] /Contents 4 0 R >>"
        ).encode(),
    ]
    content = f"q {image.width} 0 0 {image.height} 0 0 cm /Im0 Do Q".encode()
    objects.append(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
    objects.append(
        (
            f"<< /Type /XObject /Subtype /Image /Width {image.width} /Height {image.height} "
            f"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Length {len(raw_rgb)} >>\nstream\n"
        ).encode()
        + raw_rgb
        + b"\nendstream"
    )

    buf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for idx, obj_body in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += f"{idx} 0 obj\n".encode() + obj_body + b"\nendobj\n"
    xref_offset = len(buf)
    total_objs = len(objects) + 1
    buf += f"xref\n0 {total_objs}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets[1:]:
        buf += f"{off:010d} 00000 n \n".encode()
    buf += (
        f"trailer\n<< /Size {total_objs} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
    ).encode()
    return bytes(buf)


def build_mixed_pdf(pages: list[str | list[str]]) -> bytes:
    """A single multi-page PDF mixing real embedded-text pages and real scanned-image pages.

    Each entry in `pages` is either a `str` (a text-stream page, exactly
    like `build_minimal_pdf`'s non-`None` entries) or a `list[str]` (an
    image-XObject page rendered via Pillow, exactly like
    `build_scanned_pdf_page`, one page-worth of `lines`) -- so a single
    document can exercise "trusted text extraction for these pages, real
    OCR for those pages" in one real PDF, not two separate ones.
    """
    n_pages = len(pages)
    catalog_num = 1
    pages_num = 2
    first_page_num = 3
    page_nums = list(range(first_page_num, first_page_num + n_pages))
    next_free = first_page_num + n_pages

    page_objects: list[bytes] = []
    resource_objects: list[bytes] = []
    content_objects: list[bytes] = []
    page_resource_refs: list[int] = []
    page_content_refs: list[int] = []
    page_media_box: list[str] = []

    for spec in pages:
        if isinstance(spec, str):
            font_num = next_free
            next_free += 1
            content_num = next_free
            next_free += 1
            escaped = spec.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_body = f"BT /F1 12 Tf 72 700 Td ({escaped}) Tj ET".encode()
            resource_objects.append(
                (font_num, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
            )
            content_objects.append(
                (
                    content_num,
                    f"<< /Length {len(stream_body)} >>\nstream\n".encode()
                    + stream_body
                    + b"\nendstream",
                )
            )
            page_resource_refs.append(font_num)
            page_content_refs.append(content_num)
            page_media_box.append("[0 0 612 792]")
        else:
            image = Image.new("RGB", (900, 60 * len(spec) + 40), color=(255, 255, 255))
            draw = ImageDraw.Draw(image)
            font = _load_font(32)
            for index, line in enumerate(spec):
                draw.text((20, 20 + index * 60), line, fill=(0, 0, 0), font=font)
            raw_rgb = image.tobytes()

            image_num = next_free
            next_free += 1
            content_num = next_free
            next_free += 1
            content = f"q {image.width} 0 0 {image.height} 0 0 cm /Im0 Do Q".encode()
            resource_objects.append(
                (
                    image_num,
                    (
                        f"<< /Type /XObject /Subtype /Image /Width {image.width} "
                        f"/Height {image.height} /ColorSpace /DeviceRGB /BitsPerComponent 8 "
                        f"/Length {len(raw_rgb)} >>\nstream\n"
                    ).encode()
                    + raw_rgb
                    + b"\nendstream",
                )
            )
            content_objects.append(
                (
                    content_num,
                    f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream",
                )
            )
            page_resource_refs.append(image_num)
            page_content_refs.append(content_num)
            page_media_box.append(f"[0 0 {image.width} {image.height}]")

    for i, page_num in enumerate(page_nums):
        is_text_page = isinstance(pages[i], str)
        resource_key = "/Font << /F1" if is_text_page else "/XObject << /Im0"
        page_objects.append(
            (
                page_num,
                (
                    f"<< /Type /Page /Parent {pages_num} 0 R "
                    f"/Resources << {resource_key} {page_resource_refs[i]} 0 R >> >> "
                    f"/MediaBox {page_media_box[i]} /Contents {page_content_refs[i]} 0 R >>"
                ).encode(),
            )
        )

    all_objects: list[tuple[int, bytes]] = []
    kids = " ".join(f"{p} 0 R" for p in page_nums)
    all_objects.append((catalog_num, f"<< /Type /Catalog /Pages {pages_num} 0 R >>".encode()))
    all_objects.append((pages_num, f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode()))
    all_objects.extend(page_objects)
    all_objects.extend(content_objects)
    all_objects.extend(resource_objects)
    all_objects.sort(key=lambda pair: pair[0])

    buf = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for obj_num, obj_body in all_objects:
        offsets[obj_num] = len(buf)
        buf += f"{obj_num} 0 obj\n".encode() + obj_body + b"\nendobj\n"
    xref_offset = len(buf)
    total_objs = max(offsets) + 1
    buf += f"xref\n0 {total_objs}\n".encode() + b"0000000000 65535 f \n"
    for obj_num in range(1, total_objs):
        buf += f"{offsets[obj_num]:010d} 00000 n \n".encode()
    buf += (
        f"trailer\n<< /Size {total_objs} /Root {catalog_num} 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode()
    return bytes(buf)


def build_minimal_pdf(page_texts: list[str | None]) -> bytes:
    """A minimal, valid multi-page PDF. `None` for a page means an empty
    content stream (simulates a scanned/image-only page)."""
    n_pages = len(page_texts)
    catalog_num = 1
    pages_num = 2
    first_page_num = 3
    first_content_num = first_page_num + n_pages
    font_num = first_content_num + n_pages

    page_nums = list(range(first_page_num, first_page_num + n_pages))
    content_nums = list(range(first_content_num, first_content_num + n_pages))

    objects: list[bytes] = []
    kids = " ".join(f"{p} 0 R" for p in page_nums)
    objects.append(f"<< /Type /Catalog /Pages {pages_num} 0 R >>".encode())
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())

    for page_num in page_nums:
        content_num = content_nums[page_nums.index(page_num)]
        objects.append(
            (
                f"<< /Type /Page /Parent {pages_num} 0 R "
                f"/Resources << /Font << /F1 {font_num} 0 R >> >> "
                f"/MediaBox [0 0 612 792] /Contents {content_num} 0 R >>"
            ).encode()
        )

    for text in page_texts:
        if text:
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_body = f"BT /F1 12 Tf 72 700 Td ({escaped}) Tj ET".encode()
        else:
            stream_body = b""
        objects.append(
            f"<< /Length {len(stream_body)} >>\nstream\n".encode() + stream_body + b"\nendstream"
        )

    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    buf = bytearray()
    buf += b"%PDF-1.4\n"
    offsets = [0]
    for idx, obj_body in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += f"{idx} 0 obj\n".encode() + obj_body + b"\nendobj\n"

    xref_offset = len(buf)
    total_objs = len(objects) + 1
    buf += f"xref\n0 {total_objs}\n".encode()
    buf += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        buf += f"{off:010d} 00000 n \n".encode()
    buf += (
        f"trailer\n<< /Size {total_objs} /Root {catalog_num} 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode()

    return bytes(buf)


def build_encrypted_pdf(password: str = "secret123") -> bytes:
    """A single blank page PDF, encrypted with `pypdf`'s own writer."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt(password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def build_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """A DOCX with the given paragraphs, plus an optional table."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row_values in enumerate(table_rows):
            for c, value in enumerate(row_values):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_xlsx(
    headers: list[str],
    rows: list[list[object]],
    *,
    sheet_name: str = "Sheet1",
    formula_cell: tuple[int, int, str] | None = None,
) -> bytes:
    """An XLSX with one sheet. `formula_cell` is `(row_index, col_index, formula)`,
    0-based over `rows`, written as an actual Excel formula (e.g. "=1+1")."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if formula_cell is not None:
        row_idx, col_idx, formula = formula_cell
        # +2: row 1 is the header, openpyxl rows are 1-based.
        sheet.cell(row=row_idx + 2, column=col_idx + 1, value=formula)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()
